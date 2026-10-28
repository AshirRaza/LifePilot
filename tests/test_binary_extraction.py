"""Fixture tests for Phase 5B PDF/DOCX structured extraction."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from sqlalchemy import delete, func, select

from lifepilot.crawl.storage import write_bytes_atomic
from lifepilot.db.models import (
    Authority,
    Document,
    DocumentExtraction,
    DocumentVersion,
    Jurisdiction,
    Source,
)
from lifepilot.extract.docx_structured import (
    EXTRACTOR_NAME as DOCX_NAME,
)
from lifepilot.extract.docx_structured import (
    EXTRACTOR_VERSION as DOCX_VERSION,
)
from lifepilot.extract.docx_structured import extract_docx_document
from lifepilot.extract.html_structured import extract_html_document
from lifepilot.extract.pdf_structured import (
    EXTRACTOR_NAME as PDF_NAME,
)
from lifepilot.extract.pdf_structured import (
    EXTRACTOR_VERSION as PDF_VERSION,
)
from lifepilot.extract.pdf_structured import extract_pdf_document
from lifepilot.extract.runner import (
    STATUS_DRY_RUN,
    STATUS_FAILED,
    STATUS_SUCCESS,
    STATUS_UNCHANGED,
    extract_batch,
    extract_document_version,
    extract_html_batch,
)
from lifepilot.extract.schema import normalized_content_sha256


def _clear(session) -> None:
    session.execute(delete(DocumentExtraction))
    session.execute(delete(DocumentVersion))
    session.execute(delete(Document))
    session.execute(delete(Source))
    session.execute(delete(Authority))
    session.execute(delete(Jurisdiction))
    session.flush()


def _make_multipage_pdf() -> bytes:
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Residence permit overview", fontsize=18)
    page1.insert_text((72, 110), "This form explains eligibility requirements.", fontsize=11)
    page1.insert_text((72, 140), "• Bring a valid passport", fontsize=11)
    page1.insert_text((72, 160), "• Bring proof of income", fontsize=11)
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Application instructions", fontsize=16)
    page2.insert_text(
        (72, 110),
        "Complete every field. Field name: nationality. Field name: date of birth.",
        fontsize=11,
    )
    # Page 3: intentionally no extractable text (image-only / blank).
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def _make_formish_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Application form 7501", fontsize=16)
    page.insert_text((72, 120), "Surname / family name", fontsize=11)
    page.insert_text((72, 150), "Given names", fontsize=11)
    page.insert_text((72, 180), "Instructions: write in block letters.", fontsize=11)
    widget = pymupdf.Widget()
    widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    widget.field_name = "surname"
    widget.field_label = "Surname / family name"
    widget.rect = pymupdf.Rect(250, 110, 450, 130)
    page.add_widget(widget)
    data = doc.tobytes()
    doc.close()
    return data


def _make_docx() -> bytes:
    document = DocxDocument()
    document.core_properties.title = "Highly skilled migrant checklist"
    document.add_heading("Highly skilled migrant checklist", level=1)
    document.add_paragraph("Employers must submit the application.")
    document.add_heading("Documents", level=2)
    document.add_paragraph("Passport copy", style="List Bullet")
    document.add_paragraph("Employment contract", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Required"
    table.rows[1].cells[0].text = "MVV"
    table.rows[1].cells[1].text = "Yes"
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def _seed_binary_doc(
    session,
    tmp_path: Path,
    *,
    content_kind: str,
    payload: bytes,
    url: str,
    discovery_status: str = "discovered",
) -> tuple[Document, DocumentVersion]:
    country = session.scalar(select(Jurisdiction))
    if country is None:
        country = Jurisdiction(
            code="XX-BIN",
            name="Binaryland",
            jurisdiction_type="country",
            country_code="XX",
        )
        session.add(country)
        session.flush()
    authority = session.scalar(select(Authority))
    if authority is None:
        authority = Authority(jurisdiction_id=country.id, name="Binary Authority")
        session.add(authority)
        session.flush()
    source = session.scalar(select(Source))
    if source is None:
        source = Source(
            authority_id=authority.id,
            jurisdiction_id=country.id,
            name="Binary Source",
            seed_url="https://example.test/bin",
            seed_url_normalized="https://example.test/bin",
            review_status="approved",
        )
        session.add(source)
        session.flush()

    ext = "pdf" if content_kind == "pdf" else "docx"
    doc = Document(
        authority_id=authority.id,
        primary_source_id=source.id,
        canonical_url=url,
        url_normalized=url,
        content_kind=content_kind,
        language="en",
        discovery_status=discovery_status,
    )
    session.add(doc)
    session.flush()
    rel = f"raw/documents/{doc.id}/v1.{ext}"
    write_bytes_atomic(tmp_path / rel, payload)
    version = DocumentVersion(
        document_id=doc.id,
        fetched_via_source_id=source.id,
        version_number=1,
        source_url=url,
        content_type=(
            "application/pdf"
            if content_kind == "pdf"
            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        language="en",
        retrieved_at=datetime.now(timezone.utc),
        http_status=200,
        content_sha256="c" * 64,
        byte_size=len(payload),
        storage_path=rel,
        is_current=True,
    )
    session.add(version)
    session.flush()
    return doc, version


def test_pdf_multipage_page_numbers_and_reading_order():
    structured = extract_pdf_document(_make_multipage_pdf())
    assert structured.page_count == 3
    assert 3 in structured.empty_pages
    assert structured.blocks
    pages = [b.page for b in structured.blocks]
    assert pages == sorted(p for p in pages if p is not None)
    assert all(b.page in {1, 2} for b in structured.blocks)
    texts = " ".join(b.text or "" for b in structured.blocks if b.type != "table")
    assert "Residence permit overview" in texts
    assert texts.index("Residence permit overview") < texts.index(
        "Application instructions"
    )
    lists = [b for b in structured.blocks if b.type == "list"]
    assert lists
    assert any("passport" in item.lower() for item in (lists[0].items or ()))


def test_pdf_form_labels_and_instructions():
    structured = extract_pdf_document(_make_formish_pdf())
    blob = json.dumps(structured.to_content_dict()).lower()
    assert "surname" in blob
    assert "given names" in blob
    assert "block letters" in blob
    assert all(b.page == 1 for b in structured.blocks)


def test_pdf_empty_page_reported_without_invented_text():
    structured = extract_pdf_document(_make_multipage_pdf())
    assert structured.empty_pages == (3,)
    assert not any(b.page == 3 for b in structured.blocks)


def test_pdf_malformed_raises():
    try:
        extract_pdf_document(b"%PDF-1.4 not-a-real-pdf")
        raised = False
    except ValueError:
        raised = True
    assert raised


def test_docx_headings_paragraphs_lists_tables():
    structured = extract_docx_document(_make_docx())
    assert structured.title == "Highly skilled migrant checklist"
    types = [b.type for b in structured.blocks]
    assert "heading" in types
    assert "paragraph" in types
    assert "list" in types
    assert "table" in types
    headings = [b for b in structured.blocks if b.type == "heading"]
    assert headings[0].level == 1
    assert any(b.level == 2 for b in headings)
    lists = [b for b in structured.blocks if b.type == "list"]
    assert lists[0].items == ("Passport copy", "Employment contract")
    tables = [b for b in structured.blocks if b.type == "table"]
    assert tables[0].headers == ("Item", "Required")
    assert tables[0].rows == (("MVV", "Yes"),)
    assert all(b.page is None for b in structured.blocks)


def test_docx_malformed_raises():
    try:
        extract_docx_document(b"PK\x03\x04not-a-docx")
        raised = False
    except ValueError:
        raised = True
    assert raised


def test_pdf_hash_stable_and_idempotent_persist(db_session, tmp_path: Path):
    _clear(db_session)
    payload = _make_multipage_pdf()
    a = extract_pdf_document(payload)
    b = extract_pdf_document(payload)
    assert normalized_content_sha256(a) == normalized_content_sha256(b)

    doc, version = _seed_binary_doc(
        db_session,
        tmp_path,
        content_kind="pdf",
        payload=payload,
        url="https://example.test/a.pdf",
    )
    first = extract_document_version(
        db_session, version, doc, persist=True, storage_root=tmp_path
    )
    assert first.status == STATUS_SUCCESS
    assert first.extractor_name == PDF_NAME
    assert first.extractor_version == PDF_VERSION
    second = extract_document_version(
        db_session, version, doc, persist=True, storage_root=tmp_path
    )
    assert second.status == STATUS_UNCHANGED
    assert second.extraction_id == first.extraction_id
    assert (
        db_session.scalar(select(func.count()).select_from(DocumentExtraction)) == 1
    )
    stored = json.loads((tmp_path / first.storage_path).read_text(encoding="utf-8"))
    assert stored["extractor"] == PDF_NAME
    assert stored["document_id"] == doc.id
    assert stored["document_version_id"] == version.id
    assert stored["page_count"] == 3
    assert 3 in stored["empty_pages"]
    assert all("page" in block for block in stored["blocks"])


def test_docx_persist_and_dry_run(db_session, tmp_path: Path):
    _clear(db_session)
    payload = _make_docx()
    doc, version = _seed_binary_doc(
        db_session,
        tmp_path,
        content_kind="docx",
        payload=payload,
        url="https://example.test/a.docx",
    )
    dry = extract_document_version(
        db_session, version, doc, persist=False, storage_root=tmp_path
    )
    assert dry.status == STATUS_DRY_RUN
    assert db_session.scalar(select(func.count()).select_from(DocumentExtraction)) == 0
    assert not (tmp_path / "extracted").exists()

    persisted = extract_document_version(
        db_session, version, doc, persist=True, storage_root=tmp_path
    )
    assert persisted.status == STATUS_SUCCESS
    assert persisted.extractor_name == DOCX_NAME
    assert persisted.extractor_version == DOCX_VERSION
    row = db_session.scalar(select(DocumentExtraction))
    assert row is not None
    assert row.document_version_id == version.id
    assert row.extractor_name == DOCX_NAME


def test_missing_raw_file_failed(db_session, tmp_path: Path):
    _clear(db_session)
    doc, version = _seed_binary_doc(
        db_session,
        tmp_path,
        content_kind="pdf",
        payload=_make_multipage_pdf(),
        url="https://example.test/missing.pdf",
    )
    (tmp_path / version.storage_path).unlink()
    result = extract_document_version(
        db_session, version, doc, persist=True, storage_root=tmp_path
    )
    assert result.status == STATUS_FAILED
    row = db_session.scalar(select(DocumentExtraction))
    assert row is not None
    assert row.status == "failed"


def test_batch_pdf_only_filter(db_session, tmp_path: Path):
    _clear(db_session)
    pdf_doc, _ = _seed_binary_doc(
        db_session,
        tmp_path,
        content_kind="pdf",
        payload=_make_multipage_pdf(),
        url="https://example.test/only.pdf",
    )
    _seed_binary_doc(
        db_session,
        tmp_path,
        content_kind="docx",
        payload=_make_docx(),
        url="https://example.test/only.docx",
    )
    report = extract_batch(
        db_session,
        persist=True,
        max_documents=10,
        storage_root=tmp_path,
        kinds=["pdf"],
    )
    assert len(report.results) == 1
    assert report.results[0].document_id == pdf_doc.id
    assert report.results[0].extractor_name == PDF_NAME


def test_html_extraction_no_regression():
    html = b"""<!DOCTYPE html><html><head><title>T - GOV.UK</title></head>
    <body><main id="content"><h1>T</h1><p>Hello</p><ul><li>One</li></ul></main></body></html>"""
    doc = extract_html_document(html)
    assert doc.title == "T"
    assert [b.type for b in doc.blocks] == ["heading", "paragraph", "list"]


def test_html_batch_still_html_only(db_session, tmp_path: Path):
    _clear(db_session)
    _seed_binary_doc(
        db_session,
        tmp_path,
        content_kind="pdf",
        payload=_make_multipage_pdf(),
        url="https://example.test/htmlbatch.pdf",
    )
    report = extract_html_batch(
        db_session, persist=False, max_documents=5, storage_root=tmp_path
    )
    assert report.results == []
    assert "html_page" in report.message or "no matching" in report.message


def _make_junk_heading_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Eligibility overview", fontsize=18)
    page.insert_text((72, 110), "2. Requirements", fontsize=16)
    page.insert_text((72, 150), "2", fontsize=16)
    page.insert_text((72, 180), "\uf06f Yes", fontsize=16)
    page.insert_text((72, 210), "\uf06f Male", fontsize=16)
    page.insert_text((72, 240), "\uf06f", fontsize=16)
    # Enough body text so title sizes stay above the body estimate.
    y = 280
    for line in (
        "Sponsors must meet income rules.",
        "The employer submits supporting evidence.",
        "Keep a copy of every appendix appendix.",
        "Processing times depend on completeness.",
        "Contact IND if your details change.",
    ):
        page.insert_text((72, y), line, fontsize=11)
        y += 18
    data = doc.tobytes()
    doc.close()
    return data


def _make_split_title_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        "Application for the purpose of residence of 'highly skilled",
        fontsize=18,
    )
    page.insert_text((72, 96), "migrant' (recognised sponsor)", fontsize=18)
    page.insert_text((72, 140), "1", fontsize=16)
    page.insert_text((72, 180), "Who can submit this application?", fontsize=14)
    y = 220
    for line in (
        "Only a recognised sponsor may file this form.",
        "Read the explanation before you start.",
        "Send the cover page with the application.",
        "Use A4 format and actual size when printing.",
    ):
        page.insert_text((72, y), line, fontsize=11)
        y += 18
    data = doc.tobytes()
    doc.close()
    return data


def _make_widget_dedupe_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Application form", fontsize=16)
    page.insert_text((72, 120), "Surname / family name", fontsize=11)
    page.insert_text((72, 150), "Unique widget only label", fontsize=11)
    # Duplicate of drawn text — should be suppressed.
    w1 = pymupdf.Widget()
    w1.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    w1.field_name = "surname"
    w1.field_label = "Surname / family name"
    w1.rect = pymupdf.Rect(250, 110, 450, 130)
    page.add_widget(w1)
    # Distinct label not otherwise present as a separate extracted field value path.
    w2 = pymupdf.Widget()
    w2.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    w2.field_name = "only_widget"
    w2.field_label = "Only on acroform"
    w2.rect = pymupdf.Rect(250, 200, 450, 220)
    page.add_widget(w2)
    data = doc.tobytes()
    doc.close()
    return data


def test_pdf_rejects_junk_headings_keeps_numbered_heading():
    structured = extract_pdf_document(_make_junk_heading_pdf())
    headings = [b.text for b in structured.blocks if b.type == "heading"]
    assert any("Eligibility overview" in (t or "") for t in headings)
    assert any((t or "").startswith("2. Requirements") for t in headings)
    assert "2" not in headings
    assert not any("\uf06f" in (t or "") for t in headings)
    assert not any((t or "").strip() in {"Yes", "Male"} for t in headings)
    # Checkbox options survive as ordinary text/list content.
    blob = json.dumps(structured.to_content_dict())
    assert "Yes" in blob
    assert "Male" in blob
    assert "Sponsors must meet income rules." in blob


def test_pdf_merges_split_title_headings():
    structured = extract_pdf_document(_make_split_title_pdf())
    headings = [b.text for b in structured.blocks if b.type == "heading"]
    assert any(
        t
        and "highly skilled" in t
        and "migrant" in t
        and "recognised sponsor" in t
        for t in headings
    )
    assert "1" not in headings
    assert not any(
        t == "migrant' (recognised sponsor)" or t == "migrant’ (recognised sponsor)"
        for t in headings
    )


def test_pdf_dedupes_acroform_labels_against_page_text():
    structured = extract_pdf_document(_make_widget_dedupe_pdf())
    blob = json.dumps(structured.to_content_dict()).lower()
    assert blob.count("surname / family name") == 1
    assert "only on acroform" in blob
    assert PDF_VERSION == "1.1.0"
